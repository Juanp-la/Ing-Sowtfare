"""ExtractorTexto (Strategy), reglas del Extractor y estados de error."""
import pytest

from apps.analisis_hv import servicios
from apps.analisis_hv.agentes import AgenteExtractor
from apps.analisis_hv.extractor_texto import ErrorLecturaHV, ExtractorPDF, ExtractorWord, obtener_extractor
from apps.analisis_hv.parser_hv import parsear_hv
from core.agentes.trazador import Trazador
from core.datos_demo import PERFIL_DEMO, TEXTO_HV_DEMO
from core.llm.mock import MockLLMClient
from core.models import HojaVida
from core.pruebas.fabricas import docx_bytes, pdf_bytes

from .utilidades import subir, subir_y_analizar


def _pdf_sin_texto():
    import pymupdf

    documento = pymupdf.open()
    documento.new_page()
    return documento.tobytes()


def test_extrae_texto_de_pdf_y_word_con_el_mismo_resultado():
    texto_pdf = ExtractorPDF().extraer(pdf_bytes(TEXTO_HV_DEMO))
    texto_docx = ExtractorWord().extraer(docx_bytes(TEXTO_HV_DEMO))
    assert "Analista de soporte" in texto_pdf and "Analista de soporte" in texto_docx
    esperado = {**PERFIL_DEMO, "resumen": PERFIL_DEMO["resumen"]
                + " Busco un rol de soporte de aplicaciones o analista de datos junior."}
    assert parsear_hv(texto_pdf) == parsear_hv(texto_docx) == esperado


def test_word_con_tablas():
    from io import BytesIO

    from docx import Document

    documento = Document()
    documento.add_paragraph("HABILIDADES")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "SQL"
    tabla.rows[0].cells[1].text = "Excel"
    salida = BytesIO()
    documento.save(salida)
    assert parsear_hv(ExtractorWord().extraer(salida.getvalue()))["habilidades"] == ["SQL", "Excel"]


@pytest.mark.parametrize("contenido", [b"%PDF-1.4 roto", _pdf_sin_texto()])
def test_pdf_danado_o_escaneado_no_se_puede_leer(contenido):
    with pytest.raises(ErrorLecturaHV):
        ExtractorPDF().extraer(contenido)


def test_formato_desconocido():
    with pytest.raises(ErrorLecturaHV):
        obtener_extractor("odt")


def test_reglas_toleran_formatos_libres():
    perfil = parsear_hv("Juan Sintético\nResumen:\nDesarrollador\nExperiencia\nPracticante en laboratorio\n"
                        "Detalle de tareas\nFormación\nBachiller 2019\nIdiomas\nInglés A2; Portugués B1")
    assert perfil["contacto"]["nombre"] == "Juan Sintético"
    assert perfil["experiencia"][0]["cargo"] == "Practicante en laboratorio"
    assert perfil["experiencia"][0]["descripcion"] == "Detalle de tareas"
    assert perfil["educacion"] == [{"titulo": "Bachiller 2019", "institucion": "", "anio": "2019"}]
    assert perfil["idiomas"] == ["Inglés A2", "Portugués B1"]
    assert parsear_hv("")["contacto"]["nombre"] == ""


def test_agente_extractor_usa_reglas_y_llm():
    llm = MockLLMClient()
    salida = AgenteExtractor(llm=llm, trazador=Trazador(persistir=False))({"texto_hv": TEXTO_HV_DEMO})
    assert salida["perfil"] == parsear_hv(TEXTO_HV_DEMO) and llm.llamadas[0]["esquema"] == "PerfilHV"
    assert salida["perfil"]["experiencia"] == PERFIL_DEMO["experiencia"]


@pytest.mark.django_db
def test_hv_ilegible_muestra_mensaje_y_queda_en_error(cliente_autenticado):
    hv, respuesta = subir_y_analizar(cliente_autenticado, "escaneada.pdf", _pdf_sin_texto())
    assert "No fue posible leer el archivo, intente con otro formato" in respuesta.content.decode()
    hv.refresh_from_db()
    assert hv.estado == HojaVida.Estado.ERROR


@pytest.mark.django_db
def test_hv_sin_archivo_ni_texto_queda_en_error(usuario):
    hv = HojaVida.objects.create(usuario=usuario, formato="pdf")
    with pytest.raises(ErrorLecturaHV):
        servicios.analizar_hv(usuario, hv)


@pytest.mark.django_db
def test_error_del_grafo_marca_la_hv_en_error(usuario, cliente_autenticado):
    subir(cliente_autenticado)
    hv = HojaVida.objects.latest("id")

    class ServicioQueFalla:
        def ejecutar(self, *args, **kwargs):
            return {"error": "extractor: ValueError"}

    with pytest.raises(ErrorLecturaHV):
        servicios.analizar_hv(usuario, hv, servicio=ServicioQueFalla())
    hv.refresh_from_db()
    assert hv.estado == HojaVida.Estado.ERROR


@pytest.mark.django_db
def test_detalle_de_hv_analizada_ofrece_corregir_perfil(cliente_autenticado):
    hv, _ = subir_y_analizar(cliente_autenticado)
    html = cliente_autenticado.get(f"/hv/{hv.pk}/").content.decode()
    assert "Ver y corregir perfil" in html
    assert "Perfil confirmado" not in cliente_autenticado.get("/hv/").content.decode()
