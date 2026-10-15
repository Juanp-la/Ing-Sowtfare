"""ExtractorTexto (patrón Strategy): PDF con PyMuPDF y Word con python-docx. Sin OCR."""
import io
from abc import ABC, abstractmethod


class ErrorLecturaHV(Exception):
    """El archivo no se pudo leer o no tiene texto seleccionable (p. ej. HV escaneada)."""


class ExtractorTexto(ABC):
    def extraer(self, contenido):
        try:
            texto = self._extraer(contenido)
        except ErrorLecturaHV:
            raise
        except Exception as exc:  # archivo dañado o con otra estructura
            raise ErrorLecturaHV("archivo ilegible") from exc
        texto = "\n".join(linea.rstrip() for linea in texto.splitlines()).strip()
        if not texto:
            raise ErrorLecturaHV("sin texto seleccionable")
        return texto

    @abstractmethod
    def _extraer(self, contenido):
        """Devuelve el texto plano del documento."""


class ExtractorPDF(ExtractorTexto):
    def _extraer(self, contenido):
        import pymupdf

        with pymupdf.open(stream=contenido, filetype="pdf") as documento:
            return "\n".join(pagina.get_text("text") for pagina in documento)


class ExtractorWord(ExtractorTexto):
    def _extraer(self, contenido):
        from docx import Document

        documento = Document(io.BytesIO(contenido))
        lineas = [p.text for p in documento.paragraphs]
        for tabla in documento.tables:
            for fila in tabla.rows:
                lineas.append(" | ".join(celda.text for celda in fila.cells))
        return "\n".join(lineas)


EXTRACTORES = {"pdf": ExtractorPDF, "docx": ExtractorWord}


def obtener_extractor(formato):
    try:
        return EXTRACTORES[formato]()
    except KeyError as exc:
        raise ErrorLecturaHV(f"formato no soportado: {formato}") from exc
